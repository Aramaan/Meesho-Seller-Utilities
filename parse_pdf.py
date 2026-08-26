import re
from pathlib import Path

import polars as pl
from docx import Document
from pypdf import PdfReader


# Extra white spaces are removed, and leading/trailing whitespace is stripped.
def clean(value: str | None) -> str | None:
    if not value:
        return None
    return re.sub(r"\s+", " ", value).strip()


# Convert a string representing a monetary value to a float, removing commas.
def money(value: str) -> float:
    return float(value.replace(",", ""))


# Find the first match of a regex pattern in a string and return
#  the first capture group, cleaned.
def first_match(pattern: str, text: str, flags: int = 0) -> str | None:
    match = re.search(pattern, text, flags)
    return clean(match.group(1)) if match else None


# Folder containing the PDF files to be processed,
#  and the output folder for the extracted data.
PDF_DIR = Path("order_labels")

OUTPUT_DIR = Path("Tabular_data")
OUTPUT_DIR.mkdir(exist_ok=True)


def parse_pdf(pdf_path: Path) -> tuple:
    """Extract order, customer, product, and invoice data from one PDF."""

    # Read and extract text from the PDF file.
    reader = PdfReader(pdf_path)
    text = "\n".join(page.extract_text() or "" for page in reader.pages)

    # Extract the customer address block from the text.
    customer_match = re.search(
        r"Customer Address\s+(.*?)\s+If undelivered, return to:",
        text,
        re.IGNORECASE | re.DOTALL,
    )

    customer_block = customer_match.group(1).strip() if customer_match else ""

    customer_lines = [
        clean(line)
        for line in customer_block.splitlines()
        if clean(line)
    ]
    # Find the product section and its details.
    product_match = re.search(
        r"Product Details.*?SKU Size Qty Color Order No\.\s+"
        r"(?P<sku>\S+(?:\s+0)?)\s+"
        r"(?P<size>(?:(?:Un|Semi)\s+Stitched|Free\s+Size))\s+"
        r"(?P<quantity>\d+)\s+"
        r"(?P<color>\S+)\s+"
        r"(?P<item_order_number>\S+)",
        text,
        re.IGNORECASE | re.DOTALL,
    )

    # Find the invoice section and its details.
    invoice_match = re.search(
        r"Description Qty Gross Amount Discount Total\s+"
        r"(?P<description>.*?)\s+(?P<invoice_quantity>\d+)\s+"
        r"Rs\.(?P<gross_amount>[\d,.]+)\s+"
        r"Rs\.(?P<discount>[\d,.]+)\s+"
        r"Rs\.(?P<total_amount>[\d,.]+)",
        text,
        re.IGNORECASE | re.DOTALL,
    )

    # halt processing if either the product or invoice section is not found.
    if not product_match:
        raise ValueError(f"Product row not found in {pdf_path.name}")

    if not invoice_match:
        raise ValueError(f"Invoice row not found in {pdf_path.name}")

    product = product_match.groupdict()
    invoice = invoice_match.groupdict()

    # Return a tuple containing the extracted data, cleaned and formatted.
    return (
        pdf_path.name,                                                                                 # source_file
        customer_lines[0] if customer_lines else None,                                                 # customer_name
        clean(" ".join(customer_lines[1:])),                                                           # customer_address
        first_match(r"Customer Address.*?\b(\d{6})\b", text, re.IGNORECASE | re.DOTALL),               # customer_pincode
        first_match(r"BILL OF SUPPLY.*?Order No\.\s+([A-Za-z0-9]+)", text, re.IGNORECASE | re.DOTALL), # order_number
        first_match(r"Invoice No\.\s+([A-Za-z0-9]+)", text, re.IGNORECASE),                            # invoice_number
        first_match(r"Order Date\s+([0-9.]+)", text),                                                  # order_date
        first_match(r"Invoice Date\s+([0-9.]+)", text),                                                # invoice_date
        product["sku"],                                                                                # sku
        product["size"],                                                                               # size
        int(product["quantity"]),                                                                      # quantity
        product["color"],                                                                              # color
        product["item_order_number"],                                                                  # item_order_number
        clean(invoice["description"]),                                                                 # description
        money(invoice["gross_amount"]),                                                                # gross_amount
        money(invoice["discount"]),                                                                    # discount
        money(invoice["total_amount"]),                                                                # total_amount
    )


frames = []

for pdf_path in sorted(PDF_DIR.glob("*.pdf")):
    try:
        frames.append(parse_pdf(pdf_path))
        print(f"Parsed: {pdf_path.name}")
    except ValueError as error:
        print(f"Skipped: {error}")

# Headers for the DataFrame columns, matching the order of the tuple returned by parse_pdf.
headers = [
    "source_file", "customer_name", "customer_address", "customer_pincode",
    "order_number", "invoice_number", "order_date", "invoice_date",
    "sku", "size", "quantity", "color", "order_id",
    "description", "gross_amount", "discount", "total_amount"
]

# Create a Polars DataFrame from the list of tuples, using the defined headers.
df = (
    pl.DataFrame(frames, schema=headers, orient="row")
    if frames
    else pl.DataFrame(schema=headers)
)

print(df.select(df.columns[:-4]))


# Export the dataframe to CSV and JSON files and parquet and arrow formats
#  in the output directory.
df.write_csv(OUTPUT_DIR / "invoice_data.csv")
df.write_json(OUTPUT_DIR / "invoice_data.json")
df.write_ipc(OUTPUT_DIR / "invoice_data.arrow")
df.write_parquet(OUTPUT_DIR / "invoice_data.parquet")


def export_to_docx(dataframe: pl.DataFrame, output_path: Path) -> None:
    ''' Export the dataframe to a DOCX file with a table format. '''
    document = Document()
    document.add_heading("Invoice Data", level=1)

    table = document.add_table(rows=1, cols=len(dataframe.columns))
    table.style = "Table Grid"

    # Add dataframe column names as the table header.
    for index, column in enumerate(dataframe.columns):
        table.rows[0].cells[index].text = column

    # Add one DOCX table row for each PDF order.
    for row in dataframe.iter_rows(named=True):
        cells = table.add_row().cells

        for index, column in enumerate(dataframe.columns):
            cells[index].text = str(row[column] or "")


    # Save the completed DOCX file.
    document.save(output_path)


export_to_docx(df, OUTPUT_DIR / "invoice_data.docx")

print(f"Processed {df.height} PDF files.")

