from docx import Document
from pathlib import Path
import polars as pl
import re

# Extra white spaces are removed, and leading/trailing whitespace is stripped.
def clean(value: str | None) -> str | None:
    if not value:
        return None
    return re.sub(r"\s+", " ", value).strip()


# Convert a string representing a monetary value to a float, removing commas.
def money(value: str) -> float:
    return float(value.replace(",", ""))


# Find the first match of a regex pattern in a string and return
#  the first capture group, cleaned.
def first_match(pattern: str, text: str, flags: int = 0) -> str | None:
    match = re.search(pattern, text, flags)
    return clean(match.group(1)) if match else None



def export_to_docx(dataframe: pl.DataFrame, output_path: Path) -> None:
    ''' Export the dataframe to a DOCX file with a table format. '''
    document = Document()
    document.add_heading("Invoice Data", level=1)

    table = document.add_table(rows=1, cols=len(dataframe.columns))
    table.style = "Table Grid"

    # Add dataframe column names as the table header.
    for index, column in enumerate(dataframe.columns):
        table.rows[0].cells[index].text = column

    # Add one DOCX table row for each PDF order.
    for row in dataframe.iter_rows(named=True):
        cells = table.add_row().cells

        for index, column in enumerate(dataframe.columns):
            cells[index].text = str(row[column] or "")


    # Save the completed DOCX file.
    document.save(output_path)